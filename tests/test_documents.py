"""Document extraction and the short-lived upload store.

Spec: `docs/spec/gyrfalcon/19-tokenomics.md` §19.5.3.
"""

import io
import time

import pytest

from gyrfalcon.documents import (
    MAX_BYTES,
    extract,
    get_upload,
    purge_expired,
    store_upload,
)
from gyrfalcon.documents import extraction as extract_mod


def have(module: str) -> bool:
    try:
        __import__(module)
        return True
    except ImportError:
        return False


needs_pypdf = pytest.mark.skipif(not have("pypdf"), reason="needs `--extra documents`")
needs_docx = pytest.mark.skipif(not have("docx"), reason="needs `--extra documents`")
needs_xlsx = pytest.mark.skipif(not have("openpyxl"), reason="needs `--extra documents`")


class TestPlainText:
    def test_markdown_passes_through(self):
        result = extract("notes.md", b"# Title\n\nBody text.")
        assert "# Title" in result.text
        assert result.warnings == []

    def test_source_files_are_text(self):
        assert "def f()" in extract("a.py", b"def f():\n    pass").text

    def test_non_utf8_still_counts(self):
        """A Windows-encoded file should degrade, not fail."""
        result = extract("legacy.txt", "café".encode("latin-1"))
        assert result.text
        assert result.warnings == []

    def test_unknown_extension_is_read_but_flagged(self):
        result = extract("thing.xyz", b"content")
        assert result.text == "content"
        assert any("Unrecognised" in w for w in result.warnings)


class TestLimitsAndFailures:
    def test_oversize_file_is_refused(self):
        result = extract("big.txt", b"x" * (MAX_BYTES + 1))
        assert result.text == ""
        assert any("larger than" in w for w in result.warnings)

    def test_empty_file_is_reported(self):
        assert any("empty" in w for w in extract("e.txt", b"").warnings)

    def test_legacy_doc_gives_actionable_advice(self):
        result = extract("old.doc", b"\xd0\xcf\x11\xe0")
        assert any(".docx" in w for w in result.warnings)

    @needs_pypdf
    def test_corrupt_pdf_warns_rather_than_raising(self):
        result = extract("broken.pdf", b"%PDF-1.4 not really a pdf")
        assert result.text == ""
        assert result.warnings


@needs_pypdf
class TestPdf:
    def _blank_pdf(self, pages: int = 2) -> bytes:
        from pypdf import PdfWriter

        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        return buffer.getvalue()

    def test_image_only_pdf_says_so_instead_of_reporting_zero(self):
        """'0 tokens' would be a confident wrong answer to a real question."""
        result = extract("scan.pdf", self._blank_pdf())
        assert result.is_empty
        assert result.pages == 2
        assert any("no ocr" in w.lower() or "OCR" in w for w in result.warnings)

    def test_page_count_is_reported(self):
        assert extract("x.pdf", self._blank_pdf(3)).pages == 3

    def test_page_ceiling_truncates_and_warns(self, monkeypatch):
        monkeypatch.setattr(extract_mod, "MAX_PAGES", 2)
        result = extract("long.pdf", self._blank_pdf(5))
        assert result.truncated
        assert any("first 2 of 5" in w for w in result.warnings)


@needs_docx
class TestDocx:
    def test_paragraphs_and_tables_are_extracted(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Hello from a docx.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        buffer = io.BytesIO()
        document.save(buffer)

        result = extract("d.docx", buffer.getvalue())
        assert "Hello from a docx." in result.text
        assert "A\tB" in result.text


@needs_xlsx
class TestXlsx:
    def test_sheets_and_rows_are_extracted(self):
        import openpyxl

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Costs"
        sheet.append(["model", "usd"])
        sheet.append(["gpt", 1.5])
        buffer = io.BytesIO()
        book.save(buffer)

        result = extract("b.xlsx", buffer.getvalue())
        assert "# Costs" in result.text
        assert "gpt\t1.5" in result.text


class TestUploadStore:
    def test_roundtrip(self):
        upload = store_upload(extract("a.txt", b"hello"))
        assert get_upload(upload.id).text == "hello"

    def test_unknown_id_is_none(self):
        assert get_upload("nope") is None

    def test_ids_are_unique(self):
        a = store_upload(extract("a.txt", b"one"))
        b = store_upload(extract("b.txt", b"two"))
        assert a.id != b.id

    def test_expired_uploads_are_purged(self, monkeypatch):
        upload = store_upload(extract("a.txt", b"hello"))
        monkeypatch.setattr(extract_mod, "UPLOAD_TTL_SECONDS", 0)
        purge_expired(now=time.time() + 1)
        assert get_upload(upload.id) is None

    def test_warnings_travel_with_the_upload(self):
        upload = store_upload(extract("x.xyz", b"data"))
        assert upload.warnings
