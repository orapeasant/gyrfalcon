"""Pull text out of an uploaded file, and hold it just long enough to count it.

Spec: §19.5.3.

Two rules shape this module:

**Nothing is persisted.** Extracted text lives in an in-process dictionary with
a TTL and never touches disk, the session store, or a log line. Someone pasting
a contract in to see what it would cost to summarise has not agreed to have it
filed anywhere, and a token estimator has no reason to keep it once the number
is on screen.

**A file that yields nothing says so.** An image-only PDF extracts to an empty
string. Reporting "0 tokens" for that would be a confident, wrong answer to the
question the user actually asked, so it comes back as an explicit warning
instead. There is no OCR here.

Parsers are optional (`--extra documents`). Plain text, Markdown and source
files need nothing installed.
"""

from __future__ import annotations

import io
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from threading import Lock
from typing import Optional

from gyrfalcon.gyrfalcon_logging import get_logger

logger = get_logger("documents")

#: Ceilings, per §19.10 Q4. Generous enough for real documents, small enough
#: that a parse cannot wedge the dashboard's event loop for long.
MAX_BYTES = 5 * 1024 * 1024
MAX_PAGES = 100

#: How long an extraction stays available to estimate against.
UPLOAD_TTL_SECONDS = 3600

_TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".log", ".sql", ".html", ".htm", ".xml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".rb", ".php",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".swift", ".kt", ".sh", ".bash", ".ps1",
}

_INSTALL_HINT = "install the extra: `uv sync --extra documents`"


@dataclass
class Extracted:
    """The result of reading one file."""

    text: str
    filename: str = ""
    pages: Optional[int] = None
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def chars(self) -> int:
        return len(self.text)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


@dataclass
class Upload:
    """An extraction held in memory until it expires."""

    id: str
    text: str
    filename: str
    created_at: float
    pages: Optional[int] = None
    warnings: list[str] = field(default_factory=list)


_uploads: dict[str, Upload] = {}
_uploads_lock = Lock()


# ── Extraction ────────────────────────────────────────────────────────────────

def _decode_text(data: bytes) -> str:
    """Decode as UTF-8, falling back to latin-1 rather than failing.

    Log files in this project are UTF-8 by convention, but uploads come from
    anywhere; refusing a Windows-encoded file would be unhelpful when a lossy
    decode still yields a usable token count.
    """
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes, result: Extracted) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:
        result.warnings.append(f"PDF support is not installed. {_INSTALL_HINT}")
        return

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as err:
        result.warnings.append(f"Could not read the PDF: {err}")
        return

    total = len(reader.pages)
    result.pages = total
    pages = reader.pages[:MAX_PAGES]
    if total > MAX_PAGES:
        result.truncated = True
        result.warnings.append(
            f"Only the first {MAX_PAGES} of {total} pages were counted."
        )

    chunks = []
    for index, page in enumerate(pages):
        try:
            chunks.append(page.extract_text() or "")
        except Exception as err:                 # a single bad page, not the file
            result.warnings.append(f"Page {index + 1} could not be read: {err}")
    result.text = "\n\n".join(c for c in chunks if c)

    if result.is_empty:
        result.warnings.append(
            "No extractable text — this looks like a scanned or image-only PDF. "
            "There is no OCR, so no token count can be given."
        )


def _extract_docx(data: bytes, result: Extracted) -> None:
    try:
        import docx
    except ImportError:
        result.warnings.append(f"DOCX support is not installed. {_INSTALL_HINT}")
        return
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as err:
        result.warnings.append(f"Could not read the document: {err}")
        return

    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    result.text = "\n".join(parts)


def _extract_xlsx(data: bytes, result: Extracted) -> None:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError:
        result.warnings.append(f"XLSX support is not installed. {_INSTALL_HINT}")
        return
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                      data_only=True)
    except Exception as err:
        result.warnings.append(f"Could not read the workbook: {err}")
        return

    parts = []
    for sheet in book.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    result.text = "\n".join(parts)
    book.close()


def extract(filename: str, data: bytes) -> Extracted:
    """Extract text from an uploaded file. Never raises."""
    result = Extracted(text="", filename=filename)

    if len(data) > MAX_BYTES:
        result.warnings.append(
            f"File is larger than the {MAX_BYTES // (1024 * 1024)} MB limit."
        )
        return result
    if not data:
        result.warnings.append("The file is empty.")
        return result

    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        _extract_pdf(data, result)
    elif suffix == ".docx":
        _extract_docx(data, result)
    elif suffix in (".xlsx", ".xlsm"):
        _extract_xlsx(data, result)
    elif suffix in _TEXT_SUFFIXES or not suffix:
        result.text = _decode_text(data)
    elif suffix == ".doc":
        result.warnings.append(
            "Legacy .doc is not supported — save it as .docx and try again."
        )
    else:
        # Unknown extension: try it as text rather than refusing outright, but
        # say so, because binary decoded as latin-1 counts as nonsense tokens.
        result.text = _decode_text(data)
        result.warnings.append(
            f"Unrecognised file type {suffix!r}; read as plain text."
        )

    if result.is_empty and not result.warnings:
        result.warnings.append("No text could be extracted from this file.")
    return result


# ── The short-lived upload store ──────────────────────────────────────────────

def purge_expired(now: Optional[float] = None) -> int:
    """Drop expired uploads. Returns how many were removed."""
    cutoff = (now or time.time()) - UPLOAD_TTL_SECONDS
    with _uploads_lock:
        stale = [k for k, v in _uploads.items() if v.created_at < cutoff]
        for key in stale:
            del _uploads[key]
    return len(stale)


def store_upload(extracted: Extracted) -> Upload:
    """Hold an extraction in memory under a random id."""
    purge_expired()
    upload = Upload(
        id=uuid.uuid4().hex,
        text=extracted.text,
        filename=extracted.filename,
        created_at=time.time(),
        pages=extracted.pages,
        warnings=list(extracted.warnings),
    )
    with _uploads_lock:
        _uploads[upload.id] = upload
    # Deliberately logs the size, never the content.
    logger.debug("documents: stored upload %s (%d chars)", upload.id,
                 len(upload.text))
    return upload


def get_upload(upload_id: str) -> Optional[Upload]:
    """Return a stored upload, or None if it is unknown or expired."""
    purge_expired()
    with _uploads_lock:
        return _uploads.get(upload_id)
